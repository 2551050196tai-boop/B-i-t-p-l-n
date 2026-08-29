from __future__ import annotations

import argparse
import copy
import re
import zipfile
from pathlib import Path

from lxml import etree, html
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


WORD_CSS = r"""
@page Section1 {
  size: 595.35pt 841.95pt;
  margin: 42.52pt 42.52pt 42.52pt 56.69pt;
  mso-header-margin: 28.35pt;
  mso-footer-margin: 24pt;
  mso-paper-source: 0;
}
div.Section1 { page: Section1; }
* { box-sizing: border-box; }
body {
  font-family: "Times New Roman", serif;
  font-size: 12pt;
  line-height: 1.28;
  color: #000;
  margin: 0;
  padding: 0;
}
.page {
  page-break-after: always;
  break-after: page;
  margin: 0;
  padding: 0;
}
.page:last-child { page-break-after: auto; break-after: auto; }
.codex-page-marker {
  font-size: 1pt;
  line-height: 1pt;
  margin: 0;
  padding: 0;
  color: #ffffff;
}
.page-footer { display: none !important; }
.page-number { display: none !important; }
.word-page-04 table { font-size: 8.2pt; margin: 2.5pt 0; line-height: 1.05; }
.word-page-04 th, .word-page-04 td { padding: 1.2pt 2.5pt; line-height: 1.05; }
.word-page-04 .table-caption { font-size: 9.2pt; margin: 2pt 0 1pt 0; }
h1 {
  font-family: "Times New Roman", serif;
  font-size: 15pt;
  font-weight: bold;
  text-transform: uppercase;
  text-align: center;
  margin: 0 0 7pt 0;
}
h2 {
  font-family: "Times New Roman", serif;
  font-size: 13.5pt;
  font-weight: bold;
  margin: 8pt 0 4pt 0;
}
h3 {
  font-family: "Times New Roman", serif;
  font-size: 12.5pt;
  font-weight: bold;
  margin: 7pt 0 3pt 0;
}
h4 { font-family: "Times New Roman", serif; }
p {
  text-align: justify;
  text-indent: 28.35pt;
  margin: 2.5pt 0;
  line-height: 1.28;
}
p.no-indent { text-indent: 0; }
ul, ol { margin: 3pt 0; padding-left: 20pt; }
li { margin: 0 0 2pt 0; text-align: justify; line-height: 1.22; }
table {
  border-collapse: collapse;
  width: 100%;
  margin: 5pt 0;
  font-family: "Times New Roman", serif;
  font-size: 9.4pt;
}
table, th, td { border: 0.75pt solid #000; }
th {
  background: #f3f4f6;
  font-weight: bold;
  text-align: center;
  padding: 3pt 4pt;
}
td { padding: 2.5pt 4pt; vertical-align: middle; }
.table-caption {
  font-size: 10pt;
  font-style: italic;
  font-weight: bold;
  text-align: center;
  margin: 4pt 0 2pt 0;
}
pre {
  font-family: Consolas, "Courier New", monospace;
  font-size: 8.2pt;
  line-height: 1.05;
  background: #f8fafc;
  border: 0.75pt solid #cbd5e1;
  border-left: 2.25pt solid #2563eb;
  padding: 5pt;
  margin: 4pt 0;
  white-space: pre-wrap;
  word-break: break-all;
}
code {
  font-family: Consolas, "Courier New", monospace;
  font-size: 8.3pt;
  word-break: break-all;
}
.signature-section { margin-top: 12pt; }
.signature-space { height: 28pt; }
"""

TABLE_WIDTH_PERCENTAGES = [
    [8, 24, 17, 39, 12],
    [25, 16, 25, 34],
    [25, 49, 26],
    [16, 21, 24, 19, 8, 12],
    [88, 12],
    [18, 36, 46],
    [28, 36, 36],
    [9, 46, 30, 15],
    [28, 14, 58],
    [8, 28, 26, 22, 16],
    [24, 18, 18, 20, 20],
    [24, 20, 26, 14, 16],
    [25, 44, 13, 18],
    [24, 28, 48],
    [25, 14, 21, 20, 20],
    [8, 32, 15, 25, 20],
    [8, 28, 18, 46],
    [28, 18, 18, 18, 18],
    [25, 25, 36, 14],
    [28, 18, 26, 28],
    [33, 32, 35],
    [28, 62, 10],
]


def _append_style(element, extra: str) -> None:
    current = (element.get("style") or "").strip()
    if current and not current.endswith(";"):
        current += ";"
    element.set("style", current + extra)


def prepare_html(source: Path, output: Path) -> None:
    root = html.fromstring(source.read_text(encoding="utf-8"))
    report = next(e for e in root.iter() if e.get("id") == "report-content")
    pages = [
        copy.deepcopy(child)
        for child in report
        if hasattr(child, "get") and "page" in ((child.get("class") or "").split())
    ]
    if len(pages) != 20:
        raise ValueError(f"Expected 20 report pages, found {len(pages)}")

    for page_index, page in enumerate(pages):
        classes = (page.get("class") or "").split()
        classes.append(f"word-page-{page_index + 1:02d}")
        page.set("class", " ".join(classes))
        removable = page.xpath(
            './/*[contains(concat(" ", normalize-space(@class), " "), " page-footer ") '
            'or contains(concat(" ", normalize-space(@class), " "), " page-number ")]'
        )
        for footer in list(removable):
            footer.getparent().remove(footer)
        for image in page.iter("img"):
            src = image.get("src")
            if src and not src.startswith(("data:", "http://", "https://", "file:")):
                asset = (source.parent / src).resolve()
                if asset.exists():
                    image.set("src", asset.as_uri())
        _append_style(page, "page-break-after:always;")
        if page_index == len(pages) - 1:
            _append_style(page, "page-break-after:auto;")

    # Word's HTML importer does not support flexbox. Recreate the cover's
    # vertical rhythm with explicit Word-friendly margins.
    cover_content = pages[0].xpath('.//*[contains(concat(" ", normalize-space(@class), " "), " page-content ")]')[0]
    cover_blocks = [c for c in cover_content if isinstance(c.tag, str)]
    cover_content.set("style", "text-align:center;")
    if len(cover_blocks) >= 4:
        _append_style(cover_blocks[0], "text-align:center;margin:0;")
        _append_style(cover_blocks[1], "text-align:center;margin-top:76pt;margin-bottom:0;")
        _append_style(cover_blocks[2], "text-align:left;margin-top:72pt;")
        _append_style(cover_blocks[3], "text-align:center;margin-top:12pt;")

    serialized_pages = "\n".join(
        f'<p class="codex-page-marker">[[CODEX_PAGE_{index:02d}]]</p>'
        + html.tostring(page, encoding="unicode", method="html")
        for index, page in enumerate(pages, 1)
    )
    document = f"""<!DOCTYPE html>
<html xmlns:o="urn:schemas-microsoft-com:office:office"
      xmlns:w="urn:schemas-microsoft-com:office:word"
      xmlns="http://www.w3.org/TR/REC-html40" lang="vi">
<head>
  <meta charset="utf-8">
  <meta name="ProgId" content="Word.Document">
  <meta name="Generator" content="Codex HTML to DOCX">
  <title>Báo cáo Bài tập lớn - Cooks Delight</title>
  <style>{WORD_CSS}</style>
</head>
<body><div class="Section1">{serialized_pages}</div></body>
</html>"""
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(document, encoding="utf-8")


def _set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])
    _set_run_font(run, "Times New Roman", 10.5)


def _set_footer(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), "D1D5DB")
    borders.append(top)
    prefix = paragraph.add_run("Trang ")
    _set_run_font(prefix, "Times New Roman", 10.5)
    _add_field(paragraph, "PAGE")
    separator = paragraph.add_run(" / ")
    _set_run_font(separator, "Times New Roman", 10.5)
    _add_field(paragraph, "NUMPAGES")


def _set_cell_margins(cell, top=60, start=90, bottom=60, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _fix_table_geometry(
    table, content_width_dxa: int = 9920, proportions: list[float] | None = None
) -> None:
    cols = len(table.columns)
    if not cols:
        return
    grid = table._tbl.tblGrid
    if proportions is not None and len(proportions) == cols and sum(proportions) > 0:
        existing = list(proportions)
    else:
        existing = []
        if grid is not None:
            for col in grid.gridCol_lst:
                try:
                    existing.append(int(col.w))
                except (TypeError, ValueError):
                    existing.append(0)
        if len(existing) != cols or sum(existing) <= 0:
            existing = [1] * cols
    total = sum(existing)
    widths = [max(360, round(content_width_dxa * value / total)) for value in existing]
    widths[-1] += content_width_dxa - sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(content_width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "90")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _format_table_text(
    table, font_size: float, exact_line_spacing: float | None = None
) -> None:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = (
                    Pt(exact_line_spacing) if exact_line_spacing else 1.0
                )
                for run in paragraph.runs:
                    name = run.font.name or "Times New Roman"
                    if "Courier" not in name and "Consolas" not in name:
                        name = "Times New Roman"
                    _set_run_font(run, name, font_size)


def _compact_assignment_page(document: Document) -> None:
    """Keep the dense page-4 submission/assignment tables on one A4 page."""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "THÔNG TIN NỘP BÀI VÀ PHÂN CÔNG":
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = Pt(15)
            for run in paragraph.runs:
                _set_run_font(run, "Times New Roman", 13.5)
        elif text.startswith(("Bảng 0.3:", "Bảng 0.4:")):
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = Pt(10)
            for run in paragraph.runs:
                _set_run_font(run, "Times New Roman", 9.5)


def _resize_assignment_signatures(table) -> None:
    """Normalize WPS-imported VML images so signatures fit their table cells."""
    signature_sizes = ((44, 27), (35, 26), (48, 11))
    for row, (width_pt, height_pt) in zip(table.rows[1:], signature_sizes):
        cell = row.cells[5]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
        for pict in cell._tc.xpath(".//w:pict"):
            for element in pict.iter():
                if element.tag.endswith("}shape"):
                    element.set(
                        "style", f"height:{height_pt}pt;width:{width_pt}pt;"
                    )


def _materialize_page_breaks(document: Document) -> int:
    marker_re = re.compile(r"\[\[CODEX_PAGE_(\d{2})\]\]")
    marker_paragraphs = []
    for paragraph in document.paragraphs:
        match = marker_re.fullmatch(paragraph.text.strip())
        if match:
            marker_paragraphs.append((int(match.group(1)), paragraph))
    if len(marker_paragraphs) != 20:
        raise ValueError(f"Expected 20 page markers, found {len(marker_paragraphs)}")

    for page_number, paragraph in marker_paragraphs:
        marker_node = paragraph._p
        if page_number > 1:
            following = marker_node.getnext()
            while following is not None and following.tag != qn("w:p"):
                following = following.getnext()
            if following is None:
                raise ValueError(f"No paragraph found after page marker {page_number}")
            p_pr = following.get_or_add_pPr()
            if p_pr.find(qn("w:pageBreakBefore")) is None:
                p_pr.append(OxmlElement("w:pageBreakBefore"))
        marker_node.getparent().remove(marker_node)
    return len(marker_paragraphs)


def _sanitize_wps_package(source: Path, output: Path) -> None:
    """Remove dangling WPS relationships such as image Target='../NULL'."""
    rels_name = "word/_rels/document.xml.rels"
    document_name = "word/document.xml"
    with zipfile.ZipFile(source, "r") as package:
        rels_root = etree.fromstring(package.read(rels_name))
        bad_ids = []
        for rel in list(rels_root):
            target = (rel.get("Target") or "").replace("\\", "/")
            if target.upper().endswith("/NULL") or target.upper() == "NULL":
                bad_ids.append(rel.get("Id"))
                rels_root.remove(rel)

        document_root = etree.fromstring(package.read(document_name))
        if bad_ids:
            namespaces = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                "v": "urn:schemas-microsoft-com:vml",
            }
            for rel_id in bad_ids:
                images = document_root.xpath(
                    f'.//v:imagedata[@r:id="{rel_id}"]', namespaces=namespaces
                )
                for image in images:
                    pict = image
                    while pict is not None and pict.tag != qn("w:pict"):
                        pict = pict.getparent()
                    if pict is not None and pict.getparent() is not None:
                        pict.getparent().remove(pict)

        output.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as clean:
            for info in package.infolist():
                if info.filename == rels_name:
                    data = etree.tostring(
                        rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                elif info.filename == document_name:
                    data = etree.tostring(
                        document_root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                else:
                    data = package.read(info.filename)
                clean.writestr(info, data)


def finalize_docx(raw_docx: Path, final_docx: Path) -> None:
    sanitized_docx = raw_docx.with_name(raw_docx.stem + "_sanitized.docx")
    _sanitize_wps_package(raw_docx, sanitized_docx)
    document = Document(sanitized_docx)
    _materialize_page_breaks(document)
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(20)
        section.right_margin = Mm(15)
        section.top_margin = Mm(15)
        section.bottom_margin = Mm(15)
        section.header_distance = Mm(10)
        section.footer_distance = Mm(8)
        _set_footer(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    if len(document.tables) != len(TABLE_WIDTH_PERCENTAGES):
        raise ValueError(
            f"Expected {len(TABLE_WIDTH_PERCENTAGES)} tables, found {len(document.tables)}"
        )
    for index, table in enumerate(document.tables):
        _fix_table_geometry(table, proportions=TABLE_WIDTH_PERCENTAGES[index])
        if index in (2, 3):
            for row in table.rows:
                for cell in row.cells:
                    _set_cell_margins(cell, top=25, start=45, bottom=25, end=45)
            if index == 2:
                _format_table_text(table, 8.3, exact_line_spacing=9.0)
            else:
                _format_table_text(table, 8.0, exact_line_spacing=8.8)
                _resize_assignment_signatures(table)
        else:
            _format_table_text(table, 9.4)
    _compact_assignment_page(document)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.core_properties.title = "Báo cáo Bài tập lớn - Cooks Delight"
    document.core_properties.subject = "Môn Thiết kế Web (INTE03010)"
    document.core_properties.author = "Nhóm Cooks Delight"
    document.core_properties.keywords = "Cooks Delight, Thiết kế Web, Báo cáo"
    final_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(final_docx)


def main() -> None:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    prepare = sub.add_parser("prepare-html")
    prepare.add_argument("source", type=Path)
    prepare.add_argument("output", type=Path)
    finalize = sub.add_parser("finalize-docx")
    finalize.add_argument("raw_docx", type=Path)
    finalize.add_argument("final_docx", type=Path)
    args = parser.parse_args()
    if args.command == "prepare-html":
        prepare_html(args.source, args.output)
    else:
        finalize_docx(args.raw_docx, args.final_docx)


if __name__ == "__main__":
    main()
